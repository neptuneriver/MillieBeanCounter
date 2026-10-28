"""Turn Huckleberry's history into a workbook: one master sheet, one sheet per
kind of entry, audit sheets of what the sync added from daycare, daily totals
and a legend.

The same tables are written as an Excel workbook, one CSV per sheet (for
pandas), a SQLite database (for SQL), and optionally a Google Sheet. A CSV in
Huckleberry's own export format is made too, so you can check this export
against the one Huckleberry emails you.
"""

from __future__ import annotations

import csv
import io
import json
import os
import sqlite3
import tarfile
import tempfile
from dataclasses import dataclass, field
from datetime import date, datetime, timedelta
from pathlib import Path

from .history import Entry, History, words
from .journal import Journal

OZ_TO_ML = 29.5735
HUCKLEBERRY_COLUMNS = [
    "Type",
    "Start",
    "End",
    "Duration",
    "Start Condition",
    "Start Location",
    "End Condition",
    "Notes",
    "Logged By",
]


@dataclass
class Sheet:
    name: str  # tab name, e.g. "Daycare diapers"
    table: str  # file / SQL table name, e.g. "daycare_diapers"
    columns: list[str]
    rows: list[list] = field(default_factory=list)


@dataclass
class Workbook:
    sheets: list[Sheet]
    huckleberry_rows: list[list[str]]
    generated: datetime

    def sheet(self, name: str) -> Sheet:
        return next(s for s in self.sheets if s.name == name)


# ---------------------------------------------------------------------------
# Small formatting helpers
# ---------------------------------------------------------------------------


def stamp(moment: datetime | None) -> str:
    return moment.strftime("%Y-%m-%d %H:%M") if moment else ""


def minutes(seconds: float | None) -> float | None:
    return round(seconds / 60, 1) if seconds else None


def ml(amount: float | None, units: str) -> float | None:
    if amount is None:
        return None
    return round(amount * OZ_TO_ML if units.lower() == "oz" else amount, 1)


def number(value: float | None) -> str:
    """Numbers the way Huckleberry prints them: 80 not 80.0, 4.5 stays 4.5."""
    if value is None:
        return ""
    return str(int(value)) if float(value).is_integer() else repr(float(value))


def hhmm(seconds: float | None) -> str:
    total = int((seconds or 0) // 60)
    return f"{total // 60:02d}:{total % 60:02d}"


def labels(flags: list[str]) -> str:
    """Sleep condition flags as Huckleberry writes them, e.g. 'Nursing, on own in bed'."""
    text = ", ".join(words(flag) for flag in flags)
    return text[:1].upper() + text[1:]


def joined(values) -> str:
    return ", ".join(v for v in values if v)


def plain(flag: str) -> str:
    """'longTimeToFallAsleep' -> 'long time to fall asleep', 'under_10_minutes' -> 'under 10 minutes'."""
    return words(flag).replace("_", " ")


def tidy(value):
    """80.0 -> 80 so whole numbers read cleanly; everything else unchanged."""
    if isinstance(value, float) and value.is_integer():
        return int(value)
    return value


# ---------------------------------------------------------------------------
# Sheets
# ---------------------------------------------------------------------------

LEAD = ["key", "source", "date", "start"]

TYPE_SHEETS = [
    # (entry type, tab name, table name, extra columns)
    ("diaper", "Diapers", "diapers", ["kind", "pee_size", "poo_size", "color", "consistency", "diaper_rash"]),
    ("sleep", "Sleep", "sleep", ["end", "minutes", "fell_asleep", "location", "woke_up"]),
    ("bottle", "Bottles", "bottles", ["amount", "unit", "amount_ml", "bottle_type"]),
    (
        "breastfeeding",
        "Breastfeeding",
        "breastfeeding",
        ["end", "minutes", "left_minutes", "right_minutes", "last_side"],
    ),
    (
        "pumping",
        "Pumping",
        "pumping",
        ["end", "minutes", "left_amount", "right_amount", "total_amount", "unit", "total_ml"],
    ),
    ("medicine", "Medicine", "medicine", ["medicine", "amount", "unit"]),
    ("solids", "Solids", "solids", ["foods", "reactions"]),
    ("temperature", "Temperature", "temperature", ["temperature", "unit"]),
    ("growth", "Growth", "growth", ["weight", "weight_unit", "height", "height_unit", "head", "head_unit"]),
    ("activity", "Activities", "activities", ["end", "minutes", "activity"]),
    ("potty", "Potty", "potty", ["kind", "result", "pee_size", "poo_size", "color", "consistency"]),
]
TAIL = ["notes", "huckleberry_id"]

MASTER_COLUMNS = LEAD[:2] + ["type", "date", "start", "end", "minutes", "amount", "unit", "amount_ml", "detail"] + TAIL


def _type_values(entry: Entry) -> dict:
    f = entry.fields
    end = stamp(entry.end)
    if entry.type == "diaper":
        return {
            "kind": f["kind"],
            "pee_size": f["pee_size"],
            "poo_size": f["poo_size"],
            "color": f["color"],
            "consistency": f["consistency"],
            "diaper_rash": "yes" if f["diaper_rash"] else "",
        }
    if entry.type == "potty":
        return {k: f[k] for k in ("kind", "result", "pee_size", "poo_size", "color", "consistency")}
    if entry.type == "sleep":
        return {
            "end": end,
            "minutes": minutes(entry.seconds),
            "fell_asleep": joined(plain(x) for x in f["fell_asleep"]),
            "location": joined(plain(x) for x in f["location"]),
            "woke_up": joined(plain(x) for x in f["woke_up"]),
        }
    if entry.type == "bottle":
        return {
            "amount": f["amount"],
            "unit": f["units"],
            "amount_ml": ml(f["amount"], f["units"]),
            "bottle_type": f["bottle_type"],
        }
    if entry.type == "breastfeeding":
        return {
            "end": end,
            "minutes": minutes(entry.seconds),
            "left_minutes": round(f["left_seconds"] / 60, 1),
            "right_minutes": round(f["right_seconds"] / 60, 1),
            "last_side": f["last_side"],
        }
    if entry.type == "pumping":
        return {
            "end": end,
            "minutes": minutes(entry.seconds),
            "left_amount": f["left_amount"],
            "right_amount": f["right_amount"],
            "total_amount": f["total_amount"],
            "unit": f["units"],
            "total_ml": ml(f["total_amount"], f["units"]),
        }
    if entry.type == "medicine":
        amount = f["amount"] if f["amount"] else None  # the app stores a blank amount as 0
        return {"medicine": f["name"], "amount": amount, "unit": f["units"] if amount else ""}
    if entry.type == "solids":
        return {"foods": joined(f["foods"]), "reactions": joined(f["reactions"])}
    if entry.type == "temperature":
        return {"temperature": f["amount"], "unit": f["units"]}
    if entry.type == "growth":
        return {
            "weight": f["weight"],
            "weight_unit": f["weight_units"],
            "height": f["height"],
            "height_unit": f["height_units"],
            "head": f["head"],
            "head_unit": f["head_units"],
        }
    if entry.type == "activity":
        return {"end": end, "minutes": minutes(entry.seconds), "activity": words(f["activity"])}
    return {}


def _detail(entry: Entry) -> str:
    """One readable line for the master sheet."""
    f = entry.fields
    if entry.type in ("diaper", "potty"):
        sizes = [f"{part} {f[part + '_size']}" for part in ("pee", "poo") if f[part + "_size"]]
        bits = [f["kind"] + (f" ({', '.join(sizes)})" if sizes else ""), f["color"], f["consistency"]]
        if entry.type == "potty":
            bits.insert(0, f["result"])
        if f.get("diaper_rash"):
            bits.append("diaper rash")
        return ", ".join(b for b in bits if b)
    if entry.type == "sleep":
        bits = [
            ("fell asleep: " + joined(plain(x) for x in f["fell_asleep"])) if f["fell_asleep"] else "",
            ("where: " + joined(plain(x) for x in f["location"])) if f["location"] else "",
            ("woke: " + joined(plain(x) for x in f["woke_up"])) if f["woke_up"] else "",
        ]
        return " · ".join(b for b in bits if b)
    if entry.type == "bottle":
        return f["bottle_type"]
    if entry.type == "breastfeeding":
        bits = [
            f"left {round(f['left_seconds'] / 60)} min" if f["left_seconds"] else "",
            f"right {round(f['right_seconds'] / 60)} min" if f["right_seconds"] else "",
            f"last side {f['last_side']}" if f["last_side"] else "",
        ]
        return " · ".join(b for b in bits if b)
    if entry.type == "pumping":
        if f["entry_mode"] == "leftright":
            return f"left {number(f['left_amount'])} {f['units']} · right {number(f['right_amount'])} {f['units']}"
        return ""
    if entry.type == "medicine":
        return f["name"]
    if entry.type == "solids":
        return joined([joined(f["foods"]), joined(f["reactions"])])
    if entry.type == "growth":
        bits = [
            f"weight {number(f['weight'])} {f['weight_units']}" if f["weight"] is not None else "",
            f"height {number(f['height'])} {f['height_units']}" if f["height"] is not None else "",
            f"head {number(f['head'])} {f['head_units']}" if f["head"] is not None else "",
        ]
        return " · ".join(b for b in bits if b)
    if entry.type == "activity":
        return words(f["activity"])
    return ""


def _master_amount(entry: Entry) -> tuple:
    f = entry.fields
    if entry.type == "bottle":
        return f["amount"], f["units"], ml(f["amount"], f["units"])
    if entry.type == "pumping":
        return f["total_amount"], f["units"], ml(f["total_amount"], f["units"])
    if entry.type == "medicine" and f["amount"]:
        return f["amount"], f["units"], None
    if entry.type == "temperature":
        return f["amount"], f["units"], None
    return None, "", None


def _lead(entry: Entry) -> list:
    return [entry.key, entry.source, entry.start.date().isoformat(), stamp(entry.start)]


def master_row(entry: Entry) -> list:
    amount, unit, amount_ml = _master_amount(entry)
    return [
        entry.key,
        entry.source,
        entry.type,
        entry.start.date().isoformat(),
        stamp(entry.start),
        stamp(entry.end),
        minutes(entry.seconds),
        amount,
        unit,
        amount_ml,
        _detail(entry),
        entry.notes,
        entry.hb_id,
    ]


def _local(text: str | None, tz) -> str:
    if not text:
        return ""
    try:
        return stamp(datetime.fromisoformat(text).astimezone(tz))
    except ValueError:
        return text


DAYCARE_SHEETS = [
    # (sheet kind, tab name, table name, columns between key and the audit columns)
    ("diaper", "Daycare diapers", "daycare_diapers", ["date", "time", "kind"]),
    ("nap", "Daycare naps", "daycare_naps", ["date", "start", "end", "minutes"]),
    ("bottle", "Daycare bottles", "daycare_bottles", ["date", "time", "amount", "unit", "offered", "milk"]),
]
# Tab names before 2.0. The Google Sheet copy removes these old tabs.
OLD_TAB_NAMES = ("Agent diapers", "Agent naps", "Agent bottles")
DAYCARE_TAIL = ["teacher_comment", "daycare_email", "email_received", "added_at", "in_huckleberry"]
SHEET_KIND_WORDS = {"pee": "wet", "poo": "BM", "both": "wet + BM", "dry": "dry"}


def daycare_row(kind: str, event: dict, present: bool, tz) -> list:
    details = json.loads(event.get("details") or "{}")
    start = datetime.fromisoformat(event["start"]).astimezone(tz)
    if kind == "diaper":
        middle = [start.date().isoformat(), start.strftime("%H:%M"), SHEET_KIND_WORDS.get(details.get("diaper"), "")]
    elif kind == "nap":
        end = datetime.fromisoformat(details["end"]).astimezone(tz) if details.get("end") else None
        middle = [start.date().isoformat(), stamp(start), stamp(end), details.get("minutes")]
    else:
        middle = [
            start.date().isoformat(),
            start.strftime("%H:%M"),
            details.get("amount"),
            details.get("units", ""),
            details.get("offered"),
            details.get("milk", ""),
        ]
    return [
        event.get("entry_key") or "",
        *middle,
        details.get("comment", ""),
        event.get("email_subject") or "",
        _local(event.get("email_received"), tz),
        _local(event.get("updated"), tz),
        "yes" if present else "no",
    ]


def daily_totals(entries: list[Entry]) -> list[list]:
    if not entries:
        return []
    days: dict[date, dict] = {}

    def add(day: dict, name: str, value: float = 1.0) -> None:
        day[name] = day.get(name, 0) + value

    for entry in entries:
        day = days.setdefault(entry.start.date(), {})
        f = entry.fields
        if entry.type == "bottle":
            add(day, "bottles")
            add(day, "bottle_ml", ml(f["amount"], f["units"]) or 0)
        elif entry.type == "breastfeeding":
            add(day, "breastfeeds")
            add(day, "breastfeeding_minutes", (entry.seconds or 0) / 60)
        elif entry.type == "solids":
            add(day, "solids")
        elif entry.type == "diaper":
            add(day, "diapers")
            if f["kind"] in ("pee", "both"):
                add(day, "wet_diapers")
            if f["kind"] in ("poo", "both"):
                add(day, "dirty_diapers")
        elif entry.type == "sleep":
            add(day, "sleeps")
            add(day, "sleep_minutes", (entry.seconds or 0) / 60)
        elif entry.type == "pumping":
            add(day, "pumps")
            add(day, "pumped_ml", ml(f["total_amount"], f["units"]) or 0)
        elif entry.type == "medicine":
            add(day, "medicine_doses")
    rows = []
    first, last = min(days), max(days)
    for offset in range((last - first).days + 1):
        day = first + timedelta(days=offset)
        totals = days.get(day, {})
        rows.append(
            [day.isoformat()]
            + [
                round(totals.get(name, 0), 1) if name in DAILY_MEASURES else int(totals.get(name, 0))
                for name in DAILY_COLUMNS[1:]
            ]
        )
    return rows


DAILY_MEASURES = {"bottle_ml", "breastfeeding_minutes", "sleep_minutes", "pumped_ml"}


DAILY_COLUMNS = [
    "date",
    "bottles",
    "bottle_ml",
    "breastfeeds",
    "breastfeeding_minutes",
    "solids",
    "diapers",
    "wet_diapers",
    "dirty_diapers",
    "sleeps",
    "sleep_minutes",
    "pumps",
    "pumped_ml",
    "medicine_doses",
]

LEGEND = [
    ["key", "Permanent ID for each entry: MMDDYY + source letter + type letter + number for that day."],
    ["", "Example: 102226SD1 = the first diaper the sync logged for October 22, 2026."],
    ["source letter", "S = added by the sync from a daycare daily sheet; U = entered in the Huckleberry app."],
    [
        "type letter",
        "D diaper, N nap/sleep, B bottle, F breastfeeding, P pumping, M medicine, S solids, T temperature,",
    ],
    ["", "G growth, A activity (tummy time, bath...), O potty."],
    ["", "Keys never change once given. Something logged later for an earlier day gets the next free number."],
    ["source", "sync or app, the same as the source letter in the key."],
    ["date / start / end", "Local time where the entry was recorded, as YYYY-MM-DD HH:MM."],
    ["minutes", "Length in minutes (sleep, breastfeeding, pumping, activities)."],
    ["amount_ml / total_ml", f"Volume converted to ml (1 oz = {OZ_TO_ML} ml), so oz and ml entries add up."],
    ["sizes", "Diaper pee/poo sizes are Huckleberry's small, medium and large."],
    ["huckleberry_id", "Huckleberry's own ID for the entry (tracker/entry id)."],
    ["Daycare sheets", "Every entry the sync added from a daycare sheet, with the email it came from."],
    ["in_huckleberry", "yes if the entry is still in Huckleberry at this export; no if it was deleted there."],
    ["Daily totals", "Counted by the day each entry started. wet_diapers counts wet and wet + BM;"],
    ["", "dirty_diapers counts BM and wet + BM. A night's sleep counts on the day it began."],
    ["huckleberry_format.csv", "The same entries laid out like Huckleberry's own CSV export, for checking."],
]


def build(history: History, journal: Journal, tz, generated: datetime | None = None) -> Workbook:
    """All sheets, from history whose entries already have keys (see keys.assign)."""
    entries = history.entries
    present = {entry.key for entry in entries}
    sheets = [Sheet("All entries", "all_entries", MASTER_COLUMNS, [master_row(e) for e in entries])]

    synced = journal.synced_entries()
    for kind, name, table, middle in DAYCARE_SHEETS:
        rows = [
            daycare_row(kind, event, event.get("entry_key") in present, tz) for event in synced if event["kind"] == kind
        ]
        sheets.append(Sheet(name, table, ["key", *middle, *DAYCARE_TAIL], rows))

    for entry_type, name, table, extra in TYPE_SHEETS:
        columns = [*LEAD, *extra, *TAIL]
        rows = []
        for entry in entries:
            if entry.type != entry_type:
                continue
            values = _type_values(entry)
            rows.append([*_lead(entry), *[values.get(c) for c in extra], entry.notes, entry.hb_id])
        sheets.append(Sheet(name, table, columns, rows))

    sheets.append(Sheet("Daily totals", "daily_totals", DAILY_COLUMNS, daily_totals(entries)))
    sheets.append(Sheet("Legend", "legend", ["item", "meaning"], [list(row) for row in LEGEND]))
    for sheet in sheets:
        sheet.rows = [[tidy(v) for v in row] for row in sheet.rows]
    hb_rows = [huckleberry_row(e) for e in sorted(entries, key=lambda e: e.start_ts, reverse=True)]
    return Workbook(sheets, [r for r in hb_rows if r], generated or datetime.now(tz))


# ---------------------------------------------------------------------------
# Huckleberry's own CSV layout
# ---------------------------------------------------------------------------


def diaper_condition(f: dict) -> str:
    kind = f["kind"]
    if kind == "both":
        parts = [f"{part}:{f[part + '_size']}" for part in f["quantity_order"] if f.get(part + "_size")]
        return "Both" + (", " + " ".join(parts) if parts else "")
    if kind in ("pee", "poo"):
        size = f[kind + "_size"]
        return kind.capitalize() + (f":{size}" if size else "")
    return kind.capitalize()


def huckleberry_row(entry: Entry) -> list[str] | None:
    f = entry.fields
    start = stamp(entry.start)
    has_length = bool(entry.seconds and entry.seconds > 0)
    end = stamp(entry.end) if has_length else ""
    length = hhmm(entry.seconds) if has_length else ""
    notes = entry.notes
    if entry.type == "diaper":
        rash = "Diaper rash" if f["diaper_rash"] else ""
        return ["Diaper", start, "", f["color"], f["consistency"].capitalize(), rash, diaper_condition(f), notes, ""]
    if entry.type == "potty":
        result = f["result"].capitalize()
        return ["Potty", start, "", f["color"], f["consistency"].capitalize(), result, diaper_condition(f), notes, ""]
    if entry.type == "sleep":
        return [
            "Sleep",
            start,
            end,
            length,
            labels(f["fell_asleep"]),
            labels(f["location"]),
            labels(f["woke_up"]),
            notes,
            "",
        ]
    if entry.type == "bottle":
        amount = f"{number(f['amount'])}{f['units']}" if f["amount"] is not None else ""
        return ["Feed", start, "", "", f["bottle_type"], "Bottle", amount, notes, ""]
    if entry.type == "breastfeeding":
        right = f"{hhmm(f['right_seconds'])}R" if f["right_seconds"] > 0 else ""
        left = f"{hhmm(f['left_seconds'])}L" if f["left_seconds"] > 0 else ""
        return ["Feed", start, end, length, right, "Breast", left, notes, ""]
    if entry.type == "solids":
        return ["Solids", start, "", "", joined(f["foods"]), "", joined(f["reactions"]), notes, ""]
    if entry.type == "pumping":
        total = f"{number(f['total_amount'])}{f['units']}" if f["total_amount"] is not None else ""
        return ["Pump", start, end, length, total, "", "", notes, ""]
    if entry.type == "medicine":
        amount = f"{number(f['amount'])}{f['units']}" if f["amount"] else ""
        return ["Meds", start, "", "", amount, f["name"], "", notes, ""]
    if entry.type == "temperature":
        return ["Temperature", start, "", "", f"{number(f['amount'])}{f['units']}", "", "", notes, ""]
    if entry.type == "growth":
        values = [f"{number(f[k])}{f[k + '_units']}" if f[k] is not None else "" for k in ("weight", "height", "head")]
        return ["Growth", start, "", "", *values, notes, ""]
    if entry.type == "activity":
        label = words(f["activity"])
        return [label[:1].upper() + label[1:], start, end, length, "", "", "", notes, ""]
    return None


def huckleberry_csv(rows: list[list[str]]) -> str:
    """Quoted like Huckleberry's file: text in quotes, empty fields left empty."""

    def cell(value: str) -> str:
        return "" if value in ("", None) else '"' + str(value).replace('"', '""') + '"'

    lines = [",".join(cell(c) for c in HUCKLEBERRY_COLUMNS)]
    lines += [",".join(cell(c) for c in row) for row in rows]
    return "\n".join(lines) + "\n"


# ---------------------------------------------------------------------------
# Files
# ---------------------------------------------------------------------------


def _csv_text(sheet: Sheet) -> str:
    buffer = io.StringIO()
    writer = csv.writer(buffer, lineterminator="\n")
    writer.writerow(sheet.columns)
    for row in sheet.rows:
        writer.writerow(["" if v is None else v for v in row])
    return buffer.getvalue()


def write_sqlite(workbook: Workbook, path: Path) -> None:
    if path.exists():
        path.unlink()
    db = sqlite3.connect(path)
    try:
        with db:
            for sheet in workbook.sheets:
                types = []
                for index in range(len(sheet.columns)):
                    values = [row[index] for row in sheet.rows if row[index] not in (None, "")]
                    if values and all(isinstance(v, int) and not isinstance(v, bool) for v in values):
                        types.append("INTEGER")
                    elif values and all(isinstance(v, (int, float)) and not isinstance(v, bool) for v in values):
                        types.append("REAL")
                    else:
                        types.append("TEXT")
                columns = ", ".join(
                    f'"{c}" {t}' + (" PRIMARY KEY" if c == "key" and i == 0 else "")
                    for i, (c, t) in enumerate(zip(sheet.columns, types, strict=True))
                )
                db.execute(f'CREATE TABLE "{sheet.table}" ({columns})')
                marks = ", ".join("?" for _ in sheet.columns)
                db.executemany(
                    f'INSERT INTO "{sheet.table}" VALUES ({marks})',
                    [[None if v == "" else v for v in row] for row in sheet.rows],
                )
            hb = workbook.huckleberry_rows
            db.execute(
                'CREATE TABLE "huckleberry_format" ('
                + ", ".join(f'"{c.lower().replace(" ", "_")}" TEXT' for c in HUCKLEBERRY_COLUMNS)
                + ")"
            )
            db.executemany(f'INSERT INTO "huckleberry_format" VALUES ({", ".join("?" * 9)})', hb)
    finally:
        db.close()


def write_xlsx(workbook: Workbook, path: Path) -> bool:
    """Excel workbook with every sheet. Returns False if openpyxl isn't available."""
    try:
        from openpyxl import Workbook as ExcelBook
        from openpyxl.styles import Font
        from openpyxl.utils import get_column_letter
    except ImportError:
        return False
    book = ExcelBook()
    book.remove(book.active)
    for sheet in workbook.sheets:
        tab = book.create_sheet(sheet.name)
        tab.append(sheet.columns)
        for row in sheet.rows:
            values = ["" if v is None else v for v in row]
            tab.append(values)
            if any(isinstance(v, str) and v.startswith("=") for v in values):
                for cell in tab[tab.max_row]:
                    if isinstance(cell.value, str) and cell.value.startswith("="):
                        cell.data_type = "s"  # notes and teacher comments are text, never formulas
        for cell in tab[1]:
            cell.font = Font(bold=True)
        tab.freeze_panes = "A2"
        for index, column in enumerate(sheet.columns, start=1):
            longest = max([len(str(column))] + [len(str(r[index - 1] or "")) for r in sheet.rows[:500]])
            tab.column_dimensions[get_column_letter(index)].width = min(max(10, longest + 2), 60)
    book.save(path)
    return True


README_TEXT = """MillieBeanCounter data export
==============================

Made {generated} from everything in Huckleberry, plus the sync's own log.

  millie_data.xlsx         every sheet in one workbook (Excel, Numbers, Google Sheets)
  csv/                     one CSV per sheet, for pandas: pd.read_csv("csv/all_entries.csv")
  millie.db                the same tables in SQLite, for SQL
  huckleberry_format.csv   the entries in Huckleberry's own export layout, for checking

The Legend sheet (csv/legend.csv) explains the keys and columns.

If you ran Verify.command, verify_report.txt lists every difference from
Huckleberry's own CSV.

Every entry has a permanent key: MMDDYY + S (sync) or U (app) + type letter + number,
e.g. 102226SD1 = the first diaper the sync logged for October 22, 2026.

Quick starts
  pandas:  df = pd.read_csv("csv/all_entries.csv", parse_dates=["start", "end"])
           daily = pd.read_csv("csv/daily_totals.csv", parse_dates=["date"])
  SQL:     sqlite3 millie.db "SELECT date, wet_diapers, bottle_ml FROM daily_totals ORDER BY date DESC LIMIT 7"
           sqlite3 millie.db "SELECT type, source, COUNT(*) FROM all_entries GROUP BY 1, 2"
"""


def write_all(workbook: Workbook, directory: Path) -> list[Path]:
    """Write every format into directory (replacing an earlier export there)."""
    directory.mkdir(parents=True, exist_ok=True)
    (directory / "csv").mkdir(exist_ok=True)
    written = []
    for sheet in workbook.sheets:
        path = directory / "csv" / f"{sheet.table}.csv"
        path.write_text(_csv_text(sheet), encoding="utf-8")
        written.append(path)
    hb = directory / "huckleberry_format.csv"
    hb.write_text(huckleberry_csv(workbook.huckleberry_rows), encoding="utf-8")
    written.append(hb)
    db = directory / "millie.db"
    write_sqlite(workbook, db)
    written.append(db)
    xlsx = directory / "millie_data.xlsx"
    if write_xlsx(workbook, xlsx):
        written.append(xlsx)
    readme = directory / "README.txt"
    readme.write_text(README_TEXT.format(generated=stamp(workbook.generated)), encoding="utf-8")
    written.append(readme)
    return written


def tar_bytes(directory: Path) -> bytes:
    """The export folder as a tar archive (how files get out of Docker to the Mac)."""
    buffer = io.BytesIO()
    with tarfile.open(fileobj=buffer, mode="w") as archive:
        for path in sorted(directory.rglob("*")):
            if path.is_file():
                info = archive.gettarinfo(str(path), arcname=str(path.relative_to(directory)))
                info.uid = info.gid = 0
                info.uname = info.gname = ""
                info.mode = 0o644
                with path.open("rb") as handle:
                    archive.addfile(info, handle)
    return buffer.getvalue()


def fresh_directory(base: Path) -> Path:
    base.mkdir(parents=True, exist_ok=True)
    return Path(tempfile.mkdtemp(prefix="export-", dir=base))


def replace_directory(new: Path, final: Path) -> None:
    """Swap the new export in for the old one."""
    old = final.with_name(final.name + ".old")
    if old.exists():
        _remove_tree(old)
    if final.exists():
        os.replace(final, old)
    os.replace(new, final)
    if old.exists():
        _remove_tree(old)


def _remove_tree(path: Path) -> None:
    for child in sorted(path.rglob("*"), key=lambda p: len(p.parts), reverse=True):
        child.rmdir() if child.is_dir() else child.unlink()
    path.rmdir()
