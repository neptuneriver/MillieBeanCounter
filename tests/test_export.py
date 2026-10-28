"""The data export: history reading, audit keys, sheets, files, verification
and the Google Sheet copy."""

from __future__ import annotations

import asyncio
import contextlib
import io
import json
import sqlite3
import sys
import tarfile
import tempfile
import unittest
from pathlib import Path
from unittest import mock
from zoneinfo import ZoneInfo

from fakes import FIXTURES, FakeMailbox, Outbox, at, family, fixture, make_api, settings
from rebuild import load_into

from millie_bean_counter import __main__ as cli
from millie_bean_counter import config, data, export, gsheets, keys, verify
from millie_bean_counter.history import read_all
from millie_bean_counter.huckleberry import Huckleberry
from millie_bean_counter.journal import Journal
from millie_bean_counter.sync import run_once

TZ = ZoneInfo("America/New_York")
SAMPLE = (FIXTURES / "huckleberry-sample.csv").read_text()


class Base(unittest.TestCase):
    def setUp(self):
        self._tmp = tempfile.TemporaryDirectory()
        self.data = Path(self._tmp.name)
        self.store = family()
        load_into(self.store, SAMPLE)
        self.journal = Journal(self.data / "journal.sqlite3")

    def tearDown(self):
        self.journal.close()
        self._tmp.cleanup()

    def workbook(self) -> export.Workbook:
        history = asyncio.run(read_all(self.store, "child-1", TZ))
        keys.assign(self.journal, history.entries)
        return export.build(history, self.journal, TZ)

    def rows(self, book, name):
        sheet = book.sheet(name)
        return [dict(zip(sheet.columns, row, strict=True)) for row in sheet.rows]


class HuckleberryFormatTest(Base):
    def test_export_reproduces_huckleberrys_own_csv(self):
        book = self.workbook()
        self.assertEqual(export.huckleberry_csv(book.huckleberry_rows), SAMPLE)
        result = verify.compare(verify.read_their_csv(SAMPLE), book.huckleberry_rows)
        self.assertTrue(result.perfect)
        self.assertEqual(result.same, 22)

    def test_verify_spots_differences(self):
        book = self.workbook()
        changed = SAMPLE.replace('"90ml"', '"95ml"').replace('"Meds","2026-09-03 09:15",,,,"Vitamin D",,,\n', "")
        result = verify.compare(verify.read_their_csv(changed), book.huckleberry_rows)
        self.assertFalse(result.perfect)
        self.assertEqual(len(result.different), 1)
        self.assertEqual(len(result.only_ours), 1)
        text = verify.report(result)
        self.assertIn("End Condition (1)", text)
        self.assertIn("only in this export", text)

    def test_verify_ignores_order_of_diaper_parts_and_newer_entries(self):
        book = self.workbook()
        swapped = SAMPLE.replace("Both, poo:large pee:small", "Both, pee:small poo:large")
        older = "\n".join(line for line in swapped.splitlines() if "2026-09-03 21:25" not in line) + "\n"
        result = verify.compare(verify.read_their_csv(older), book.huckleberry_rows)
        self.assertTrue(result.perfect)
        self.assertEqual(result.ours_newer, 1)

    def test_not_a_huckleberry_file(self):
        with self.assertRaises(verify.NotAHuckleberryFile):
            verify.read_their_csv("date,amount\n2026-01-01,4\n")


class KeysTest(Base):
    def test_keys_follow_the_format_and_time_order(self):
        book = self.workbook()
        first_day = [(r["key"], r["type"]) for r in self.rows(book, "All entries") if r["date"] == "2026-09-01"]
        self.assertEqual(
            first_day,
            [
                ("090126UM1", "medicine"),
                ("090126UB1", "bottle"),
                ("090126UD1", "diaper"),
                ("090126UN1", "sleep"),
                ("090126UF1", "breastfeeding"),
            ],
        )

    def test_keys_never_change(self):
        before = {r["huckleberry_id"]: r["key"] for r in self.rows(self.workbook(), "All entries")}
        # A diaper logged later for an earlier time that day gets the next number.
        self.store.add_interval("diaper", {"mode": "pee", "start": at("2026-09-01T07:00").timestamp(), "offset": 240})
        after = self.rows(self.workbook(), "All entries")
        for row in after:
            if row["huckleberry_id"] in before:
                self.assertEqual(row["key"], before[row["huckleberry_id"]])
        new = [r for r in after if r["huckleberry_id"] not in before]
        self.assertEqual([(r["key"], r["start"]) for r in new], [("090126UD2", "2026-09-01 07:00")])

    def test_key_follows_an_entry_the_app_moves_into_a_batch(self):
        book = self.workbook()
        diaper = next(r for r in self.rows(book, "All entries") if r["key"] == "090126UD1")
        path = next(p for p in self.store.docs if len(p) == 4 and f"{p[0]}/{p[3]}" == diaper["huckleberry_id"])
        raw = self.store.docs.pop(path)
        self.store.docs[("diaper", "child-1", "intervals", "batch-1")] = {"multi": True, "data": {"other-id": raw}}
        moved = next(r for r in self.rows(self.workbook(), "All entries") if r["huckleberry_id"] == "diaper/other-id")
        self.assertEqual(moved["key"], "090126UD1")


class SheetsTest(Base):
    def test_every_sheet_is_there(self):
        book = self.workbook()
        self.assertEqual(
            [s.name for s in book.sheets],
            [
                "All entries",
                "Daycare diapers",
                "Daycare naps",
                "Daycare bottles",
                "Diapers",
                "Sleep",
                "Bottles",
                "Breastfeeding",
                "Pumping",
                "Medicine",
                "Solids",
                "Temperature",
                "Growth",
                "Activities",
                "Potty",
                "Daily totals",
                "Legend",
            ],
        )
        counts = {s.name: len(s.rows) for s in book.sheets}
        self.assertEqual(counts["All entries"], 22)
        self.assertEqual(
            (counts["Diapers"], counts["Sleep"], counts["Bottles"], counts["Breastfeeding"], counts["Pumping"]),
            (5, 3, 3, 5, 2),
        )
        self.assertEqual((counts["Medicine"], counts["Activities"], counts["Daily totals"]), (2, 2, 3))

    def test_type_sheets_have_readable_values(self):
        book = self.workbook()
        bottle = self.rows(book, "Bottles")[-1]
        self.assertEqual((bottle["amount"], bottle["unit"], bottle["amount_ml"]), (4.5, "oz", 133.1))
        self.assertEqual(bottle["notes"], 'Fed by grandma, "all of it"')
        sleep = self.rows(book, "Sleep")[-1]
        self.assertEqual(sleep["fell_asleep"], "long time to fall asleep")
        self.assertEqual(sleep["location"], "nursing, on own in bed")
        self.assertEqual(sleep["minutes"], 582)
        nap = self.rows(book, "Sleep")[0]
        self.assertEqual(nap["fell_asleep"], "10-20 minutes, happy")
        diaper = self.rows(book, "Diapers")[-1]
        self.assertEqual(
            (diaper["kind"], diaper["pee_size"], diaper["poo_size"], diaper["color"], diaper["diaper_rash"]),
            ("both", "small", "large", "green", "yes"),
        )
        feed = self.rows(book, "Breastfeeding")[-1]
        self.assertEqual((feed["left_minutes"], feed["right_minutes"]), (13.2, 12.2))

    def test_daily_totals(self):
        book = self.workbook()
        day = next(r for r in self.rows(book, "Daily totals") if r["date"] == "2026-09-03")
        self.assertEqual(
            (day["bottles"], day["bottle_ml"], day["breastfeeds"], day["diapers"], day["wet_diapers"]),
            (2, 223.1, 2, 2, 2),
        )
        self.assertEqual((day["dirty_diapers"], day["sleeps"], day["pumps"], day["pumped_ml"]), (1, 1, 1, 150))
        self.assertEqual(day["medicine_doses"], 1)


class OddEntriesTest(unittest.TestCase):
    def test_a_corrupt_entry_is_skipped_not_fatal(self):
        from millie_bean_counter.history import normalize

        self.assertIsNone(normalize("diaper", "a", {"start": 1e30, "mode": "pee"}, TZ))
        entry = normalize("diaper", "b", {"start": 1760000000, "mode": "pee", "offset": 99999}, TZ)
        self.assertEqual(entry.tz, TZ)  # nonsense offset: the configured time zone is used
        self.assertEqual(entry.start.date().isoformat(), "2025-10-09")


class FilesTest(Base):
    def test_all_formats_are_written(self):
        book = self.workbook()
        folder = self.data / "out"
        written = {str(p.relative_to(folder)) for p in export.write_all(book, folder)}
        self.assertIn("csv/all_entries.csv", written)
        self.assertIn("csv/daycare_naps.csv", written)
        self.assertIn("huckleberry_format.csv", written)
        self.assertIn("millie.db", written)
        self.assertEqual((folder / "huckleberry_format.csv").read_text(), SAMPLE)
        db = sqlite3.connect(folder / "millie.db")
        try:
            wet = db.execute("SELECT wet_diapers FROM daily_totals WHERE date = '2026-09-03'").fetchone()[0]
            self.assertEqual(wet, 2)
            schema = db.execute("SELECT sql FROM sqlite_master WHERE name = 'all_entries'").fetchone()[0]
            self.assertIn('"key" TEXT PRIMARY KEY', schema)
            self.assertEqual(db.execute("SELECT COUNT(*) FROM huckleberry_format").fetchone()[0], 22)
        finally:
            db.close()
        try:
            import openpyxl
        except ImportError:
            return
        self.assertIn("millie_data.xlsx", written)
        excel = openpyxl.load_workbook(folder / "millie_data.xlsx")
        self.assertEqual(excel.sheetnames[0], "All entries")
        self.assertEqual(excel["All entries"]["A1"].value, "key")
        self.assertEqual(excel["All entries"].max_row, 23)

    def test_text_that_looks_like_a_formula_stays_text_in_excel(self):
        try:
            import openpyxl
        except ImportError:
            self.skipTest("openpyxl isn't installed")
        book = self.workbook()
        book.sheet("All entries").rows[0][11] = '=HYPERLINK("https://example.com","click")'
        path = self.data / "book.xlsx"
        export.write_xlsx(book, path)
        cell = openpyxl.load_workbook(path)["All entries"]["L2"]
        self.assertEqual(cell.data_type, "s")
        self.assertEqual(cell.value, '=HYPERLINK("https://example.com","click")')

    def test_tar_carries_the_folder(self):
        folder = self.data / "out"
        export.write_all(self.workbook(), folder)
        with tarfile.open(fileobj=io.BytesIO(export.tar_bytes(folder))) as archive:
            names = archive.getnames()
        self.assertIn("csv/all_entries.csv", names)
        self.assertIn("README.txt", names)


class DaycareSheetsTest(unittest.TestCase):
    """A real daycare import followed by an export."""

    def test_synced_entries_get_s_keys_and_their_own_sheets(self):
        with tempfile.TemporaryDirectory() as tmp:
            store = family()
            load_into(store, "\n".join(SAMPLE.splitlines()[:1]) + "\n")  # header only: empty history
            s = settings(Path(tmp))
            factory = lambda s: Huckleberry(s, api=make_api(store))
            run_once(
                s,
                now=at("2026-10-22T22:00:00"),
                mailbox_factory=FakeMailbox([fixture("forwarded-oct-12.eml")]),
                huckleberry_factory=factory,
                notifier=Outbox(),
            )
            # Someone deletes one of the synced diapers in the app afterwards.
            gone = next(p for p, d in store.docs.items() if p[0] == "diaper" and len(p) == 4 and "intervals" in p)
            del store.docs[gone]
            result = data.export_now(s, huckleberry_factory=factory, push_sheet=False)
            book = result.workbook
            diapers = [
                dict(zip(book.sheet("Daycare diapers").columns, r, strict=True))
                for r in book.sheet("Daycare diapers").rows
            ]
            self.assertEqual([d["key"] for d in diapers], ["101226SD1", "101226SD2", "101226SD3", "101226SD4"])
            self.assertEqual([d["kind"] for d in diapers], ["wet", "wet + BM", "wet", "BM"])
            self.assertEqual(sorted(d["in_huckleberry"] for d in diapers), ["no", "yes", "yes", "yes"])
            self.assertEqual(diapers[0]["daycare_email"], "Fwd: Josephine's Daily Sheet for Monday, October 12th")
            naps = book.sheet("Daycare naps").rows
            self.assertEqual([r[0] for r in naps], ["101226SN1", "101226SN2", "101226SN3"])
            bottles = [
                dict(zip(book.sheet("Daycare bottles").columns, r, strict=True))
                for r in book.sheet("Daycare bottles").rows
            ]
            self.assertEqual(
                [(b["key"], b["amount"], b["offered"], b["teacher_comment"]) for b in bottles],
                [("101226SB1", 4, 4.5, ""), ("101226SB2", 3.5, 4.5, "Took a little break halfway through.")],
            )
            sources = {r[0]: r[1] for r in book.sheet("All entries").rows}
            self.assertEqual(sources["101226SN1"], "sync")
            self.assertEqual(len(book.sheet("All entries").rows), 8)  # 9 added, 1 deleted in the app
            self.assertTrue((Path(tmp) / "exports" / "latest" / "csv" / "daycare_bottles.csv").exists())


class FakeResponse:
    def __init__(self, status: int, body: dict | None = None):
        self.status_code = status
        self._body = body or {}
        self.content = json.dumps(self._body).encode()
        self.text = json.dumps(self._body)

    def json(self):
        return self._body


class FakeSheetsSession:
    def __init__(self, tabs=("Sheet1",), status: int = 200, message: str = ""):
        self.calls = []
        self.tabs = {name: i for i, name in enumerate(tabs)}
        self.status, self.message = status, message

    def get(self, url, params=None, timeout=None):
        self.calls.append(("get", url, params))
        if self.status >= 400:
            return FakeResponse(self.status, {"error": {"message": self.message}})
        if params and params.get("fields") == "properties.title":
            return FakeResponse(200, {"properties": {"title": "MillieBeanCounter data"}})
        sheets = [{"properties": {"sheetId": i, "title": t}} for t, i in self.tabs.items()]
        return FakeResponse(200, {"sheets": sheets})

    def post(self, url, json=None, timeout=None):
        self.calls.append(("post", url, json))
        replies = []
        for request in (json or {}).get("requests", []):
            if "addSheet" in request:
                title = request["addSheet"]["properties"]["title"]
                self.tabs[title] = len(self.tabs) + 100
                replies.append({"addSheet": {"properties": {"sheetId": self.tabs[title], "title": title}}})
        return FakeResponse(200, {"replies": replies})


class GoogleSheetTest(Base):
    def test_push_creates_tabs_and_writes_everything(self):
        book = self.workbook()
        session = FakeSheetsSession(tabs=("Sheet1", "Diapers"))
        sheet = gsheets.GoogleSheet("1AbCdEfGhIjKlMnOpQrStUvWxYz0123456789", session, "robot@example.iam")
        cells = sheet.push(book)
        posts = [c for c in session.calls if c[0] == "post"]
        added = [r["addSheet"]["properties"]["title"] for r in posts[0][2]["requests"]]
        self.assertNotIn("Diapers", added)
        self.assertIn("All entries", added)
        self.assertIn(":batchUpdate", posts[1][1])
        self.assertTrue(posts[2][1].endswith("/values:batchClear"))
        self.assertNotIn("'Sheet1'", posts[2][2]["ranges"])  # tabs you made yourself are left alone
        data_calls = [c for c in posts if c[1].endswith("/values:batchUpdate")]
        ranges = [d["range"] for call in data_calls for d in call[2]["data"]]
        self.assertIn("'All entries'!A1:M23", ranges)
        self.assertEqual(data_calls[0][2]["valueInputOption"], "RAW")
        self.assertEqual(cells, sum((len(s.rows) + 1) * len(s.columns) for s in book.sheets))

    def test_tabs_from_before_the_rename_are_removed(self):
        session = FakeSheetsSession(tabs=("Sheet1", "Agent diapers", "Agent naps", "Agent bottles"))
        gsheets.GoogleSheet("1AbCdEfGhIjKlMnOpQrStUvWxYz0123456789", session).push(self.workbook())
        posts = [c for c in session.calls if c[0] == "post" and c[1].endswith(":batchUpdate")]
        deleted = [r["deleteSheet"]["sheetId"] for r in posts[1][2]["requests"] if "deleteSheet" in r]
        self.assertEqual(deleted, [1, 2, 3])
        added = [r["addSheet"]["properties"]["title"] for r in posts[0][2]["requests"]]
        self.assertIn("Daycare diapers", added)

    def test_big_sheets_go_in_chunks(self):
        book = self.workbook()
        big = book.sheet("All entries")
        big.rows = big.rows * 400  # 8,800 rows x 13 columns
        session = FakeSheetsSession()
        gsheets.GoogleSheet("1AbCdEfGhIjKlMnOpQrStUvWxYz0123456789", session).push(book)
        data_calls = [c for c in session.calls if c[0] == "post" and c[1].endswith("/values:batchUpdate")]
        self.assertGreater(len(data_calls), 2)
        for call in data_calls:
            self.assertLessEqual(sum(len(d["values"]) * len(d["values"][0]) for d in call[2]["data"]), 40_000)

    def test_problems_are_explained(self):
        cases = [
            (403, "The caller does not have permission", "isn't shared with robot@example.iam"),
            (403, "Google Sheets API has not been used in project 1 before or it is disabled", "turned off"),
            (404, "Requested entity was not found.", "wasn't found"),
        ]
        for status, message, expected in cases:
            session = FakeSheetsSession(status=status, message=message)
            sheet = gsheets.GoogleSheet("1AbCdEfGhIjKlMnOpQrStUvWxYz0123456789", session, "robot@example.iam")
            with self.assertRaises(gsheets.SheetsError) as caught:
                sheet.title()
            self.assertIn(expected, str(caught.exception))

    def test_links_and_keys(self):
        link = "https://docs.google.com/spreadsheets/d/1AbCdEfGhIjKlMnOpQrStUvWxYz0123456789/edit#gid=0"
        self.assertEqual(gsheets.spreadsheet_id(link), "1AbCdEfGhIjKlMnOpQrStUvWxYz0123456789")
        with self.assertRaises(gsheets.SheetsError):
            gsheets.spreadsheet_id("https://example.com/nope")
        key = {"type": "service_account", "client_email": "robot@x.iam", "private_key": "k", "token_uri": "t"}
        self.assertEqual(gsheets.save_key(self.data, json.dumps(key)), "robot@x.iam")
        self.assertEqual(gsheets.load_key(self.data)["client_email"], "robot@x.iam")
        with self.assertRaises(gsheets.SheetsError):
            gsheets.save_key(self.data, json.dumps({"type": "authorized_user"}))
        gsheets.key_path(self.data).write_text("{not json")
        with self.assertRaises(gsheets.SheetsError):
            gsheets.load_key(self.data)

    def test_a_google_problem_never_loses_the_files(self):
        s = settings(self.data, GOOGLE_SHEET="1AbCdEfGhIjKlMnOpQrStUvWxYz0123456789")
        gsheets.key_path(self.data).write_text("{not json")  # a damaged key
        factory = lambda s: Huckleberry(s, api=make_api(self.store))
        result = data.export_now(s, huckleberry_factory=factory)
        self.assertIn("couldn't be read", result.sheet_error)
        self.assertTrue((self.data / "exports" / "latest" / "csv" / "all_entries.csv").exists())


class CommandLineTest(Base):
    """What Get Data.command and Verify.command run inside Docker."""

    def run_cli(self, *args: str, stdin: bytes = b"") -> tuple[int, bytes, str]:
        values = {
            "GMAIL_ADDRESS": "daycare-inbox@example.com",
            "GMAIL_APP_PASSWORD": "abcdefghijklmnop",
            "HUCKLEBERRY_EMAIL": "parent@example.com",
            "HUCKLEBERRY_PASSWORD": "x",
            "DAYCARE_CHILD": "Josephine",
        }
        config.save(values, self.data)
        fake_in = io.TextIOWrapper(io.BytesIO(stdin), encoding="utf-8")
        fake_out = io.TextIOWrapper(io.BytesIO(), encoding="utf-8")
        messages = io.StringIO()
        with contextlib.ExitStack() as stack:
            stack.enter_context(
                mock.patch.object(data, "Huckleberry", lambda s: Huckleberry(s, api=make_api(self.store)))
            )
            stack.enter_context(mock.patch.object(sys, "stdin", fake_in))
            stack.enter_context(mock.patch.object(sys, "stdout", fake_out))
            stack.enter_context(mock.patch.object(sys, "stderr", messages))
            code = cli.main(["--data", str(self.data), "export", *args])
        return code, fake_out.buffer.getvalue(), messages.getvalue()

    def test_verify_command_gets_a_tar_and_a_report(self):
        code, tar, messages = self.run_cli("--tar", "--compare", stdin=b"\xef\xbb\xbf" + SAMPLE.encode())
        self.assertEqual(code, 0, messages)
        self.assertIn("All 22 entries match", messages)
        with tarfile.open(fileobj=io.BytesIO(tar)) as archive:
            names = archive.getnames()
            report = archive.extractfile("verify_report.txt").read().decode()
        self.assertIn("csv/all_entries.csv", names)
        self.assertIn("All 22 entries match", report)

    def test_differences_give_exit_code_1_but_still_the_files(self):
        changed = SAMPLE.replace('"90ml"', '"95ml"').encode()
        code, tar, messages = self.run_cli("--tar", "--compare", stdin=changed)
        self.assertEqual(code, 1)
        self.assertIn("1 entry differs", messages)
        self.assertTrue(tar)

    def test_wrong_file_is_explained(self):
        code, tar, messages = self.run_cli("--tar", "--compare", stdin=b"name,age\nx,1\n")
        self.assertEqual(code, 1)
        self.assertEqual(tar, b"")
        self.assertIn("doesn't look like Huckleberry's CSV export", messages)


if __name__ == "__main__":
    unittest.main()
